"""Read Excel and CSV uploads into rows of cells."""

from __future__ import annotations

import csv
import io
from typing import Any

import xlrd
from openpyxl import load_workbook

from app.errors import UserError
from app.fa import display_fa, is_blank, preview_cell, suggest_columns, to_persian_digits

MAX_ROWS = 100_000
MAX_COLS = 200


def read_tabular(filename: str, data: bytes) -> dict[str, list[list[Any]]]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix in {"xlsx", "xlsm"}:
            return _read_xlsx(data)
        if suffix == "xls":
            return _read_xls(data)
        if suffix == "csv":
            return {"کاربرگ": _read_csv(data)}
    except UserError:
        raise
    except Exception as exc:  # pragma: no cover - defensive user message
        raise UserError(
            "این فایل خوانده نشد. اکسل یا CSV سالم، بدون رمز، بارگذاری کنید."
        ) from exc
    raise UserError("فقط فایل xlsx، xlsm، xls یا csv پذیرفته می‌شود.")


def _trim_trailing_empty(rows: list[list[Any]]) -> list[list[Any]]:
    while rows and all(is_blank(cell) for cell in rows[-1]):
        rows.pop()
    return rows


def _read_xlsx(data: bytes) -> dict[str, list[list[Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        message = str(exc).lower()
        if "password" in message or "encrypted" in message:
            raise UserError("این فایل رمز دارد و قابل خواندن نیست.") from exc
        raise UserError("فایل اکسل باز نشد. آن را با پسوند xlsx ذخیره کنید.") from exc
    sheets: dict[str, list[list[Any]]] = {}
    try:
        for worksheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    raise UserError(
                        f"برگه «{worksheet.title}» بیش از {to_persian_digits(MAX_ROWS)} ردیف دارد."
                    )
                values = list(row[:MAX_COLS])
                rows.append(values)
            sheets[worksheet.title] = _trim_trailing_empty(rows)
    finally:
        workbook.close()
    if not sheets:
        raise UserError("این فایل برگه‌ای ندارد.")
    return sheets


def _read_xls(data: bytes) -> dict[str, list[list[Any]]]:
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise UserError("فایل xls باز نشد.") from exc
    sheets: dict[str, list[list[Any]]] = {}
    for sheet in book.sheets():
        if sheet.nrows > MAX_ROWS:
            raise UserError(
                f"برگه «{sheet.name}» بیش از {to_persian_digits(MAX_ROWS)} ردیف دارد."
            )
        rows = []
        for row_index in range(sheet.nrows):
            row = []
            for col_index in range(min(sheet.ncols, MAX_COLS)):
                cell = sheet.cell(row_index, col_index)
                row.append(_xls_value(cell))
            rows.append(row)
        sheets[sheet.name] = _trim_trailing_empty(rows)
    if not sheets:
        raise UserError("این فایل برگه‌ای ندارد.")
    return sheets


def _xls_value(cell: xlrd.sheet.Cell) -> Any:
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
        return None
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        if cell.value == int(cell.value):
            return int(cell.value)
        return cell.value
    return cell.value


def _read_csv(data: bytes) -> list[list[Any]]:
    text = _decode_csv(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[Any]] = []
    for index, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if index >= MAX_ROWS:
            raise UserError(f"این CSV بیش از {to_persian_digits(MAX_ROWS)} ردیف دارد.")
        rows.append(row[:MAX_COLS])
    return _trim_trailing_empty(rows)


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1256"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UserError("رمزگذاری این CSV شناخته نشد. فایل را با UTF-8 ذخیره کنید.")


def make_headers(row: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    last_real = -1
    for index, cell in enumerate(row):
        raw = "" if is_blank(cell) else display_fa(preview_cell(cell))
        if raw:
            last_real = index
            text = raw
        else:
            text = f"ستون {to_persian_digits(index + 1)}"
        count = seen.get(text, 0) + 1
        seen[text] = count
        if count > 1:
            text = f"{text} ({to_persian_digits(count)})"
        headers.append(text)
    if last_real == -1:
        return []
    return headers[: last_real + 1]


def detect_header_row(rows: list[list[Any]]) -> int:
    best_row = 1
    best_score = -1
    for index, row in enumerate(rows[:25]):
        headers = make_headers(row)
        if not headers:
            continue
        suggestions = suggest_columns(headers)
        score = sum(value is not None for value in suggestions.values()) * 10
        score += min(sum(not is_blank(cell) for cell in row), 9)
        if score > best_score:
            best_score = score
            best_row = index + 1
    return best_row


def describe_sheet(rows: list[list[Any]], header_row: int) -> dict:
    if header_row < 1:
        raise UserError("ردیف عنوان باید از ۱ به بعد باشد.")
    if rows and header_row > len(rows):
        raise UserError("ردیف عنوان خارج از محدودهٔ این برگه است.")
    header_index = header_row - 1
    header_cells = rows[header_index] if rows else []
    headers = make_headers(header_cells)
    data = rows[header_index + 1 :] if rows else []
    preview = []
    data_rows = 0
    for row in data:
        if all(is_blank(cell) for cell in row):
            continue
        data_rows += 1
        if len(preview) < 8:
            shown = [preview_cell(cell) for cell in row[: len(headers)]]
            if len(shown) < len(headers):
                shown.extend([""] * (len(headers) - len(shown)))
            preview.append(shown)
    return {
        "header_row": header_row,
        "headers": headers,
        "preview": preview,
        "row_count": data_rows,
        "suggestions": suggest_columns(headers),
        "preview_limited": len(headers) > 12,
    }
