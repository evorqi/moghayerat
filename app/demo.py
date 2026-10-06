"""Sample workbooks so the page can be tried before the real files arrive."""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from app.fa import with_check_digit

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "samples"


def sample_people() -> dict[str, dict]:
    ali = with_check_digit("049281736")
    sara = with_check_digit("238471905")
    reza = with_check_digit("512670348")
    javad = with_check_digit("670198234")
    fatemeh = with_check_digit("135792468")
    hossein = with_check_digit("802345671")
    leila_due = Decimal("4000000")
    return {
        "ali": {
            "first": "عل" + "\u064a",
            "last": "رضایی",
            "nid": ali,
            "due": Decimal("20000000"),
            "bank": Decimal("12000000"),
            "fund": Decimal("8000000"),
        },
        "sara": {
            "first": "سارا",
            "last": "محمدی",
            "nid": sara,
            "due": Decimal("15000000"),
            "bank": Decimal("5000000"),
        },
        "reza": {
            "first": "رضا",
            "last": "کاظمی",
            "nid": reza,
            "due": Decimal("8000000"),
        },
        "javad": {
            "first": "جواد",
            "last": "اکبری",
            "nid": javad,
            "due": Decimal("6000000"),
            "bank": Decimal("7000000"),
        },
        "fatemeh": {
            "first": "فاطمه",
            "last": "نوری",
            "nid": fatemeh,
            "due_text": "۹,۰۰۰,۰۰۰ ریال",
            "bank": Decimal("9000000"),
        },
        "hossein": {
            "first": "حسین",
            "last": "مرادی",
            "nid": hossein,
            "due_rows": (Decimal("1000000"), Decimal("1500000")),
            "bank": Decimal("2500000"),
        },
        "leila": {
            "first": "لیلا",
            "last": "کرمی",
            "due": leila_due,
            "bank": Decimal("1000000"),
        },
        "maryam": {
            "first": "مریم",
            "last": "حسینی",
            "nid": with_check_digit("990123456"),
            "fund": Decimal("2500000"),
            "note": "کمک هزینه",
        },
    }


def write_sample_files(directory: Path | None = None) -> list[Path]:
    directory = directory or SAMPLE_DIR
    directory.mkdir(parents=True, exist_ok=True)
    people = sample_people()
    due_path = directory / "هزینه-آموزش.xlsx"
    bank_path = directory / "پرداخت-بانک.xlsx"
    fund_path = directory / "پرداخت-صندوق.csv"

    due = Workbook()
    sheet = due.active
    sheet.title = "هزینه"
    sheet.sheet_view.rightToLeft = True
    sheet["A1"] = "گزارش هزینه آموزش"
    sheet["A2"] = "سال تحصیلی ۱۴۰۴"
    headers = ["نام", "نام خانوادگی", "کد ملی", "هزینه"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(3, column, header)
    rows = [
        [people["ali"]["first"], people["ali"]["last"], int(people["ali"]["nid"]), int(people["ali"]["due"])],
        [people["sara"]["first"], people["sara"]["last"], people["sara"]["nid"], int(people["sara"]["due"])],
        [people["reza"]["first"], people["reza"]["last"], people["reza"]["nid"], int(people["reza"]["due"])],
        [people["javad"]["first"], people["javad"]["last"], people["javad"]["nid"], int(people["javad"]["due"])],
        [people["fatemeh"]["first"], people["fatemeh"]["last"], _persian_digits(people["fatemeh"]["nid"]), people["fatemeh"]["due_text"]],
        [people["hossein"]["first"], people["hossein"]["last"], people["hossein"]["nid"], int(people["hossein"]["due_rows"][0])],
        [people["hossein"]["first"], people["hossein"]["last"], people["hossein"]["nid"], int(people["hossein"]["due_rows"][1])],
        [people["leila"]["first"], people["leila"]["last"], None, int(people["leila"]["due"])],
        ["کامران", "تستی", "1234567890", "ندارد"],
        [None, None, None, None],
    ]
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(4 + offset, column, value)
            if column == 3 and isinstance(value, str):
                cell.number_format = "@"
    due.save(due_path)

    bank = Workbook()
    bank_sheet = bank.active
    bank_sheet.title = "واریز"
    bank_sheet.sheet_view.rightToLeft = True
    for column, header in enumerate(["نام و نام خانوادگی", "کد ملی", "مبلغ واریز"], start=1):
        bank_sheet.cell(1, column, header)
    bank_rows = [
        ["علی رضائی", people["ali"]["nid"], int(people["ali"]["bank"])],
        ["سارا محمدی", people["sara"]["nid"], int(people["sara"]["bank"])],
        ["جواد اکبری", people["javad"]["nid"], int(people["javad"]["bank"])],
        ["فاطمه نوری", people["fatemeh"]["nid"], int(people["fatemeh"]["bank"])],
        ["حسین مرادی", people["hossein"]["nid"], int(people["hossein"]["bank"])],
        ["لیلا کرمی", None, int(people["leila"]["bank"])],
    ]
    for offset, row in enumerate(bank_rows):
        for column, value in enumerate(row, start=1):
            cell = bank_sheet.cell(2 + offset, column, value)
            if column == 2 and isinstance(value, str):
                cell.number_format = "@"
    bank.save(bank_path)

    with fund_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["کد ملی", "نام", "نام خانوادگی", "مبلغ", "توضیحات"])
        writer.writerow([people["ali"]["nid"], "علی", "رضایی", int(people["ali"]["fund"]), "قسط دوم"])
        writer.writerow(
            [people["maryam"]["nid"], people["maryam"]["first"], people["maryam"]["last"], int(people["maryam"]["fund"]), people["maryam"]["note"]]
        )

    return [due_path, bank_path, fund_path]


def _persian_digits(value: str) -> str:
    return str(value).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))


if __name__ == "__main__":
    for path in write_sample_files():
        print(path)
