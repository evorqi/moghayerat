import io
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.demo import sample_people, write_sample_files
from app.excel_export import build_workbook
from app.fa import gregorian_to_jalali, national_id_is_valid, parse_amount, parse_month, with_check_digit
from app.main import app
from app.readers import read_tabular
from app.reconcile import (
    STATUS_OVER,
    STATUS_PAID_ONLY,
    STATUS_SETTLED,
    STATUS_SHORT,
    STATUS_UNPAID,
    ColumnMap,
    FileJob,
    reconcile,
)
from app.state import reset

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def clean_state():
    reset()
    yield
    reset()


def test_jalali_known_dates():
    assert gregorian_to_jalali(2024, 3, 20) == (1403, 1, 1)
    assert gregorian_to_jalali(2025, 3, 21) == (1404, 1, 1)
    assert gregorian_to_jalali(2026, 3, 21) == (1405, 1, 1)
    assert gregorian_to_jalali(2026, 10, 6) == (1405, 7, 14)


def test_amount_and_national_id_parsing():
    assert parse_amount("۹,۰۰۰,۰۰۰ ریال") == Decimal("9000000")
    assert parse_amount("1.200.000") == Decimal("1200000")
    assert parse_amount("(2,500)") == Decimal("-2500")
    assert parse_amount(None) is None
    code = with_check_digit("049281736")
    assert national_id_is_valid(code)
    assert not national_id_is_valid("1234567890")


def test_sample_reconciliation_totals_and_statuses(tmp_path):
    paths = write_sample_files(tmp_path)
    jobs = []
    expected_headers = {
        "هزینه-آموزش.xlsx": 3,
        "پرداخت-بانک.xlsx": 1,
        "پرداخت-صندوق.csv": 1,
    }
    for path in paths:
        sheets = read_tabular(path.name, path.read_bytes())
        sheet_name, rows = next(iter(sheets.items()))
        from app.readers import describe_sheet, detect_header_row

        header_row = detect_header_row(rows)
        assert header_row == expected_headers[path.name]
        described = describe_sheet(rows, header_row)
        suggestions = described["suggestions"]
        role = "due" if "هزینه" in path.name else "paid"
        jobs.append(
            FileJob(
                file_id=path.name,
                filename=path.name,
                sheet=sheet_name,
                header_row=header_row,
                label=path.stem,
                role=role,
                factor=Decimal("1"),
                columns=ColumnMap(
                    national_id=suggestions["national_id"],
                    first_name=suggestions["first_name"],
                    last_name=suggestions["last_name"],
                    full_name=suggestions["full_name"],
                    amount=suggestions["amount"],
                    note=suggestions["note"],
                ),
                rows=rows,
            )
        )

    report = reconcile(jobs, Decimal("0"))
    people = sample_people()
    by_id = {person.national_id: person for person in report.people if person.national_id}
    by_name = {person.full_display: person for person in report.people}

    assert len(report.people) == 8
    assert sum((person.due for person in report.people), Decimal("0")) == Decimal("64500000")
    assert sum((person.paid for person in report.people), Decimal("0")) == Decimal("47000000")
    assert sum((person.balance for person in report.people), Decimal("0")) == Decimal("17500000")

    ali = by_id[people["ali"]["nid"]]
    assert ali.due == Decimal("20000000")
    assert ali.paid == Decimal("20000000")
    assert ali.status == STATUS_SETTLED
    assert ali.first_name == "علی"
    assert "\u06cc" in ali.first_name
    assert "\u064a" not in ali.first_name
    assert ali.last_name == "رضایی"
    assert "قسط دوم" in "؛ ".join(ali.notes)
    assert "یکسان نیست" not in "؛ ".join(ali.notes)
    assert ali.national_id.startswith("0")

    sara = by_id[people["sara"]["nid"]]
    assert sara.status == STATUS_SHORT
    assert sara.balance == Decimal("10000000")
    assert by_id[people["reza"]["nid"]].status == STATUS_UNPAID
    assert by_id[people["javad"]["nid"]].status == STATUS_OVER
    assert by_id[people["javad"]["nid"]].balance == Decimal("-1000000")

    fatemeh = by_id[people["fatemeh"]["nid"]]
    assert fatemeh.first_name == "فاطمه"
    assert fatemeh.status == STATUS_SETTLED
    assert fatemeh.due == Decimal("9000000")

    hossein = by_id[people["hossein"]["nid"]]
    assert hossein.due == Decimal("2500000")
    assert hossein.paid == Decimal("2500000")
    assert hossein.status == STATUS_SETTLED

    leila = by_name["لیلا کرمی"]
    assert leila.status == STATUS_SHORT
    assert leila.balance == Decimal("3000000")
    assert any("با نام تطبیق شد" in note for note in leila.notes)

    maryam = by_id[people["maryam"]["nid"]]
    assert maryam.status == STATUS_PAID_ONLY
    assert maryam.paid == Decimal("2500000")
    assert any("کمک هزینه" in note for note in maryam.notes)
    assert any("نامعتبر" in warning for warning in report.warnings)

    workbook = build_workbook(report)
    check_workbook(workbook, ali.national_id)


def test_factor_tolerance_and_both_columns():
    rows = [
        ["نام", "کد ملی", "هزینه", "پرداخت"],
        ["علی رضایی", "1111111111", "1000", "250"],
        ["سارا محمدی", "2222222222", "نامعلوم", "80"],
    ]
    job = FileJob(
        file_id="both",
        filename="both.xlsx",
        sheet="Sheet",
        header_row=1,
        label="کلاس",
        role="both",
        factor=Decimal("10"),
        columns=ColumnMap(full_name=0, national_id=1, amount_due=2, amount_paid=3),
        rows=rows,
    )
    report = reconcile([job], Decimal("0"))
    ali = next(person for person in report.people if person.first_name == "علی")
    sara = next(person for person in report.people if person.first_name == "سارا")
    assert ali.due == Decimal("10000")
    assert ali.paid == Decimal("2500")
    assert ali.status == STATUS_SHORT
    assert any("رقم کنترلی" in note for note in ali.notes)
    assert sara.due == Decimal("0")
    assert sara.paid == Decimal("800")
    assert sara.status == STATUS_PAID_ONLY
    assert any("نامعتبر" in warning for warning in report.warnings)

    settled = reconcile([job], Decimal("8000"))
    ali = next(person for person in settled.people if person.first_name == "علی")
    assert ali.status == STATUS_SETTLED
    assert any("چشم‌پوشی" in note for note in ali.notes)


def test_cp1256_csv_roundtrip():
    # ویندوز فارسی اغلب «ی» عربی را در cp1256 ذخیره می‌کند.
    yeh = "\u064a"
    text = f"نام,مبلغ\nنگ{yeh}ن مراد{yeh},1500\n"
    rows = read_tabular("paid.csv", text.encode("cp1256"))["کاربرگ"]
    assert rows[0][0] == "نام"
    assert rows[1][0] == f"نگ{yeh}ن مراد{yeh}"


def check_workbook(payload: bytes, national_id: str):
    workbook = load_workbook(io.BytesIO(payload))
    assert workbook.sheetnames == ["خلاصه", "تطبیق و مانده", "مانده ماه‌ها", "جزئیات منابع"]
    for name in workbook.sheetnames:
        assert workbook[name].sheet_view.rightToLeft is True
    detail = workbook["تطبیق و مانده"]
    found = False
    heights = []
    for row in detail.iter_rows(min_row=4, max_row=detail.max_row - 1, max_col=8, values_only=False):
        heights.append(detail.row_dimensions[row[0].row].height)
        if row[3].value == national_id:
            found = True
            assert row[3].data_type == "s"
            assert row[1].value == "علی رضایی"
            assert row[1].alignment.readingOrder == 2
            assert not row[1].alignment.wrap_text
    assert found
    assert heights
    assert len(set(heights)) == 1
    assert any(isinstance(cell.value, str) and str(cell.value).startswith("=SUBTOTAL") for cell in detail[detail.max_row])

    archive = zipfile.ZipFile(io.BytesIO(payload))
    sheets_xml = "\n".join(
        archive.read(name).decode("utf-8")
        for name in archive.namelist()
        if name.startswith("xl/worksheets/")
    )
    assert 'rightToLeft="1"' in sheets_xml
    assert "فاطمه" in sheets_xml
    assert "همطاف" not in sheets_xml
    assert not any("\ufb50" <= char <= "\ufdff" or "\ufe70" <= char <= "\ufeef" for char in sheets_xml)
    fonts = archive.read("xl/styles.xml").decode("utf-8")
    assert "Tahoma" in fonts
    assert 'val="178"' in fonts


def test_http_demo_download():
    client = TestClient(app)
    loaded = client.post("/api/demo")
    assert loaded.status_code == 200
    files = loaded.json()["files"]
    assert len(files) == 3
    due = next(item for item in files if "هزینه" in item["filename"])
    assert due["sheets"][0]["header_row"] == 3
    body = {"tolerance": "0", "files": []}
    for item in files:
        sheet = next(sheet for sheet in item["sheets"] if sheet["name"] == item["active_sheet"])
        body["files"].append(
            {
                "id": item["id"],
                "sheet": sheet["name"],
                "header_row": sheet["header_row"],
                "label": item["label"],
                "role": item["role"],
                "factor": "1",
                "columns": sheet["suggestions"],
            }
        )
    created = client.post("/api/reconcile", json=body)
    assert created.status_code == 200, created.text
    summary = created.json()["summary"]
    assert summary["people"] == 8
    assert Decimal(summary["balance_total"]) == Decimal("17500000")
    downloaded = client.get(created.json()["download_url"])
    assert downloaded.status_code == 200
    assert downloaded.content[:2] == b"PK"
    check_workbook(downloaded.content, sample_people()["ali"]["nid"])


def test_only_shahrivar_paid_rest_remains():
    due_rows = [
        ["نام بیمه‌گذار", "شماره بیمه", "نوع بیمه", "ماه", "مبلغ"],
        ["لیلا احمدی", "B-100", "عمر", "فروردین", 1000],
        ["لیلا احمدی", "B-100", "عمر", "اردیبهشت", 1000],
        ["لیلا احمدی", "B-100", "عمر", "شهریور", 1000],
    ]
    paid_rows = [
        ["نام بیمه‌گذار", "شماره بیمه", "نوع بیمه", "ماه", "مبلغ"],
        ["لیلا احمدی", "B-100", "عمر", "شهریور", 1000],
    ]
    columns = ColumnMap(policyholder=0, insurance_no=1, insurance_type=2, month=3, amount=4)
    report = reconcile(
        [
            FileJob("due", "due.xlsx", "برگ", 1, "حق بیمه", "due", Decimal("1"), columns, due_rows),
            FileJob("paid", "paid.xlsx", "برگ", 1, "واریز", "paid", Decimal("1"), columns, paid_rows),
        ],
        Decimal("0"),
    )
    person = report.people[0]
    assert person.holder_name == "لیلا احمدی"
    assert person.insurance_no == "B-100"
    assert person.insurance_type == "عمر"
    assert person.due == Decimal("3000")
    assert person.paid == Decimal("1000")
    assert person.balance == Decimal("2000")
    assert person.paid_month_list() == ["شهریور"]
    assert person.unpaid_month_list() == ["فروردین", "اردیبهشت"]
    assert person.discrepancy_text() == "فقط شهریور پرداخت شده و بقیه مانده است"
    workbook = load_workbook(io.BytesIO(build_workbook(report)))
    month_sheet = workbook["مانده ماه‌ها"]
    months = [row[4].value for row in month_sheet.iter_rows(min_row=4, max_row=5)]
    assert months == ["فروردین", "اردیبهشت"]
    assert parse_month("6") == "شهریور"
    assert parse_month("shahrivar") == "شهریور"


def test_missing_role_is_persian_error():
    rows = [["نام", "مبلغ"], ["علی", "10"]]
    job = FileJob(
        file_id="a",
        filename="a.xlsx",
        sheet="Sheet",
        header_row=1,
        label="فقط هزینه",
        role="due",
        factor=Decimal("1"),
        columns=ColumnMap(first_name=0, amount=1),
        rows=rows,
    )
    with pytest.raises(Exception) as caught:
        reconcile([job], Decimal("0"))
    assert "پرداخت‌شده" in str(caught.value)
